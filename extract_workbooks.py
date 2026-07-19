import os, sys, json, re
from datetime import datetime
sys.path.insert(0, os.path.abspath('.temp_xlrd'))
import xlrd
import openpyxl

BASE = os.path.abspath('planilhas')
OUT = os.path.abspath('extracted_inventory.json')

def norm(v):
    s = str(v or '').strip().lower()
    s = re.sub(r'\s+', ' ', s)
    return s

def clean(v):
    if v is None or v == '': return None
    if isinstance(v, datetime): return v.isoformat()
    if isinstance(v, float) and v.is_integer(): return int(v)
    if not isinstance(v, (str, int, float, bool, list, dict)):
        return str(v)
    return v

def cell_record(file_name, sheet, address, value, formula=None, row_context=''):
    return {
        'arquivo': file_name, 'aba': sheet, 'celula': address,
        'valor': clean(value), 'formula': formula or '',
        'contexto_linha': row_context[:500],
        'tipo': 'formula' if formula else ('numero' if isinstance(value,(int,float)) else 'texto')
    }

def xls_records(path):
    wb = xlrd.open_workbook(path, formatting_info=False)
    cells=[]; fields=[]
    for sh in wb.sheets():
        rows=[]
        for r in range(sh.nrows):
            vals=[clean(sh.cell_value(r,c)) for c in range(sh.ncols)]
            rows.append(vals)
        for r, vals in enumerate(rows):
            ctx=' | '.join(str(v) for v in vals if v not in (None,''))
            for c, v in enumerate(vals):
                if v not in (None,''):
                    addr=f'{xlrd.formula.colname(c)}{r+1}'
                    cells.append(cell_record(os.path.relpath(path, BASE), sh.name, addr, v, row_context=ctx))
        header_candidates=[]
        for r, vals in enumerate(rows):
            text_count=sum(1 for v in vals if isinstance(v,str) and v.strip())
            non=sum(v not in (None,'') for v in vals)
            if non>=2 and text_count>=2:
                header_candidates.append(r)
        hr=header_candidates[0] if header_candidates else None
        if hr is not None:
            for c, v in enumerate(rows[hr]):
                if v not in (None,''):
                    col=xlrd.formula.colname(c)
                    data=[row[c] for row in rows[hr+1:] if c < len(row) and row[c] not in (None,'')]
                    fields.append({'arquivo':os.path.relpath(path,BASE),'aba':sh.name,'coluna':col,'nome_campo':str(v),'linha_cabecalho':hr+1,'qtd_preenchidos':len(data),'exemplos_valores':json.dumps(data[:5],ensure_ascii=False),'formulas':''})
    return cells, fields

def xlsx_records(path):
    wf=openpyxl.load_workbook(path,data_only=False,read_only=False)
    wd=openpyxl.load_workbook(path,data_only=True,read_only=False)
    cells=[]; fields=[]
    for ws in wf.worksheets:
        wsd=wd[ws.title]
        rows=[[c.value for c in row] for row in ws.iter_rows()]
        for r, row in enumerate(rows):
            ctx=' | '.join(str(v) for v in row if v not in (None,''))
            for c, v in enumerate(row):
                if v not in (None,''):
                    addr=ws.cell(r+1,c+1).coordinate
                    f=v if isinstance(v,str) and v.startswith('=') else ''
                    display=wsd.cell(r+1,c+1).value
                    cells.append(cell_record(os.path.relpath(path,BASE),ws.title,addr,display,f,ctx))
        header_candidates=[]
        for r,row in enumerate(rows):
            text_count=sum(1 for v in row if isinstance(v,str) and v.strip())
            non=sum(v not in (None,'') for v in row)
            if non>=2 and text_count>=2: header_candidates.append(r)
        hr=header_candidates[0] if header_candidates else None
        if hr is not None:
            for c,v in enumerate(rows[hr]):
                if v not in (None,''):
                    data=[rr[c] for rr in rows[hr+1:] if c<len(rr) and rr[c] not in (None,'')]
                    fs=[]
                    for rr in range(hr+1,len(rows)):
                        if c < len(rows[rr]) and isinstance(rows[rr][c],str) and rows[rr][c].startswith('='): fs.append(rows[rr][c])
                    fields.append({'arquivo':os.path.relpath(path,BASE),'aba':ws.title,'coluna':ws.cell(hr+1,c+1).column_letter,'nome_campo':str(v),'linha_cabecalho':hr+1,'qtd_preenchidos':len(data),'exemplos_valores':json.dumps([clean(x) for x in data[:5]],ensure_ascii=False),'formulas':' | '.join(fs[:5])})
    return cells, fields

all_cells=[]; all_fields=[]
for root,_,files in os.walk(BASE):
    for fn in sorted(files):
        if fn.lower().endswith('.xls'):
            c,f=xls_records(os.path.join(root,fn))
        elif fn.lower().endswith('.xlsx'):
            c,f=xlsx_records(os.path.join(root,fn))
        else: continue
        all_cells.extend(c); all_fields.extend(f)

by_name={}
for f in all_fields:
    by_name.setdefault(norm(f['nome_campo']),[]).append(f'{f["arquivo"]}::{f["aba"]}!{f["coluna"]}')
for f in all_fields:
    matches=by_name.get(norm(f['nome_campo']),[])
    f['correspondencias']=' | '.join(x for x in matches if x != f'{f["arquivo"]}::{f["aba"]}!{f["coluna"]}')
for c in all_cells:
    refs=re.findall(r"(?:'[^']+'|[A-Za-z_][^! ]*)!?\$?[A-Z]{1,3}\$?\d+",c['formula'])
    c['referencias_formula']=' | '.join(refs)

result={'gerado_em':datetime.now().isoformat(),'arquivos':sorted({c['arquivo'] for c in all_cells}),'campos':all_fields,'celulas':all_cells}
with open(OUT,'w',encoding='utf-8') as fp: json.dump(result,fp,ensure_ascii=False)
print(json.dumps({'arquivos':len(result['arquivos']),'campos':len(all_fields),'celulas':len(all_cells),'out':OUT},ensure_ascii=False))
